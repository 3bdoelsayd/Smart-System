import openpyxl

for filepath, level_name in [(r'C:\Users\Disney\Downloads\الفرقة_الثانية.xlsx', 'الفرقة الثانية'), (r'C:\Users\Disney\Downloads\الفرقة_الثالثة.xlsx', 'الفرقة الثالثة')]:
    print(f"=== File: {level_name} ===")
    wb = openpyxl.load_workbook(filepath)
    print("Sheets:", wb.sheetnames)
    for sheetname in wb.sheetnames:
        sheet = wb[sheetname]
        valid_rows = 0
        for i in range(2, sheet.max_row + 1):
            row = [cell.value for cell in sheet[i]]
            if any(row):
                valid_rows += 1
        print(f"  Sheet '{sheetname}': {valid_rows} students")
        if sheet.max_row >= 2:
            print(f"    Sample row 2: {[cell.value for cell in sheet[2]]}")
