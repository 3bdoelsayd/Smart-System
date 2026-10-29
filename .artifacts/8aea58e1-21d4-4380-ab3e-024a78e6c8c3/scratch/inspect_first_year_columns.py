import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\Disney\Downloads\الفرقة_الأولى.xlsx')
for sheetname in wb.sheetnames:
    sheet = wb[sheetname]
    print(f"=== Sheet: {sheetname} ===")
    headers = [cell.value for cell in sheet[1]]
    print("Headers:", headers)
    # Check sample row
    if sheet.max_row >= 2:
        print("Row 2:", [cell.value for cell in sheet[2]])
