import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\Disney\Downloads\الفرقة_الرابعة.xlsx')
print("Sheets:", wb.sheetnames)
for sheetname in wb.sheetnames:
    sheet = wb[sheetname]
    valid_rows = 0
    for i in range(2, sheet.max_row + 1):
        row = [cell.value for cell in sheet[i]]
        if any(row):
            valid_rows += 1
    print(f"Sheet '{sheetname}': {valid_rows} rows")
    if sheet.max_row >= 2:
        print(f"  Sample row 2: {[cell.value for cell in sheet[2]]}")
