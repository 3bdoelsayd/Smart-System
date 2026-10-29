import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\Disney\Downloads\الفرقة_الأولى.xlsx')
for sheetname in wb.sheetnames:
    sheet = wb[sheetname]
    print(f"--- Sheet: {sheetname} ---")
    print("Max rows:", sheet.max_row)
    rows_data = []
    for i in range(2, sheet.max_row + 1):
        vals = [cell.value for cell in sheet[i]]
        if any(vals):
            rows_data.append(vals)
    print(f"Valid student rows count: {len(rows_data)}")
    if rows_data:
        print("Sample first row:", rows_data[0])
        print("Sample last row:", rows_data[-1])
