import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\Disney\Downloads\الفرقة_الأولى.xlsx')
print("Sheet names:", wb.sheetnames)
sheet = wb.active
print("Active sheet:", sheet.title)
print("Max row:", sheet.max_row)
for i in range(1, min(15, sheet.max_row + 1)):
    row_vals = [cell.value for cell in sheet[i]]
    print(f"Row {i}: {row_vals}")
